"""Build local demo DOCX and a matching reflowable reading preview."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[2]
data = json.loads((Path(__file__).parent / 'content.json').read_text())
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.8)
for name, size in [('Normal', 12), ('Title', 24), ('Heading 1', 16)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'PingFang SC')
    style.paragraph_format.space_after = Pt(12)
    style.paragraph_format.line_spacing = 1.45
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.add_heading(data['title'], 0)
doc.add_paragraph('GAIP Agent  /  演示文件')
doc.add_paragraph(data['intro'])
for i, item in enumerate(data['sections']):
    doc.add_heading(item['title'], 1)
    for text in item['paragraphs']: doc.add_paragraph(text)
footer = section.footer.paragraphs[0]
footer.text = '演示资料  ·  '
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.save(root / 'demo-files/family-plan.docx')
# Preview is derived from the saved document, so it cannot drift from its body text.
saved = Document(root / 'demo-files/family-plan.docx')
blocks = [{'kind': 'h1' if p.style.name == 'Title' else 'h2' if p.style.name.startswith('Heading') else 'p', 'text': p.text} for p in saved.paragraphs if p.text.strip()]
(root / 'demo-files/word-preview.json').write_text(json.dumps(blocks, ensure_ascii=False))
